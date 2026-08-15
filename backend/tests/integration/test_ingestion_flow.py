"""End-to-end ingestion verification (design §8, FR-C6, FR-C7).

Runs the full pipeline against a real database and the real embedding model,
asserting the Phase 3 acceptance criterion: uploading a multi-page PDF results
in chunks with embeddings visible in the database, and the status transitions
correctly through the lifecycle.
"""

from __future__ import annotations

import io
import uuid
from datetime import UTC, datetime

import pytest
from docx import Document as DocxDocument
from sqlalchemy import func, select
from sqlalchemy.ext.asyncio import AsyncEngine, AsyncSession, async_sessionmaker

from app.core.config import Settings
from app.db.models.document import Document, DocumentStatus
from app.db.models.document_chunk import DocumentChunk
from app.db.models.project import Project
from app.db.models.user import User
from app.db.repositories.document_chunk_repository import DocumentChunkRepository
from app.db.repositories.document_repository import DocumentRepository
from app.services.embedding_service import get_embedding_service
from app.services.ingestion.pipeline import IngestionPipeline

pytestmark = pytest.mark.integration


def _multi_page_pdf_bytes() -> bytes:
    import fitz

    doc = fitz.open()
    for page_index in range(3):
        page = doc.new_page()
        body = " ".join(f"page{page_index}word{i}" for i in range(120))
        page.insert_text((72, 72), f"Page {page_index + 1} heading. {body}", fontsize=11)
    data = doc.tobytes()
    doc.close()
    return bytes(data)


def _docx_bytes() -> bytes:
    document = DocxDocument()
    document.add_heading("Introduction", level=1)
    document.add_paragraph(" ".join(f"introword{i}" for i in range(100)))
    document.add_heading("Conclusion", level=1)
    document.add_paragraph(" ".join(f"concword{i}" for i in range(100)))
    buffer = io.BytesIO()
    document.save(buffer)
    return buffer.getvalue()


async def _seed_project(session: AsyncSession) -> Project:
    user = User(
        id=uuid.uuid4(),
        email=f"ingest-{uuid.uuid4().hex}@example.com",
        hashed_password="x",
        name="Ingest User",
        created_at=datetime.now(UTC),
    )
    session.add(user)
    await session.flush()

    project = Project(
        id=uuid.uuid4(), owner_id=user.id, title="Ingest Project", created_at=datetime.now(UTC)
    )
    session.add(project)
    await session.flush()
    return project


async def _create_pending_document(
    session: AsyncSession, project: Project, *, mime_type: str, filename: str, size_bytes: int
) -> Document:
    document = Document(
        id=uuid.uuid4(),
        project_id=project.id,
        filename=filename,
        status=DocumentStatus.PENDING,
        metadata_={
            "title": None,
            "authors": [],
            "page_count": None,
            "mime_type": mime_type,
            "size_bytes": size_bytes,
            "ocr_used": False,
            "ocr_low_confidence": False,
            "ocr_min_confidence": None,
            "failure_reason": None,
        },
        uploaded_at=datetime.now(UTC),
    )
    session.add(document)
    await session.flush()
    return document


def _build_pipeline(session: AsyncSession, settings: Settings) -> IngestionPipeline:
    return IngestionPipeline(
        document_repository=DocumentRepository(session),
        chunk_repository=DocumentChunkRepository(session),
        embedding_service=get_embedding_service(settings),
        settings=settings,
    )


async def test_multi_page_pdf_produces_chunks_with_embeddings(
    db_engine: AsyncEngine, settings: Settings
) -> None:
    session_factory = async_sessionmaker(bind=db_engine, expire_on_commit=False)
    content = _multi_page_pdf_bytes()

    async with session_factory() as session:
        project = await _seed_project(session)
        document = await _create_pending_document(
            session,
            project,
            mime_type="application/pdf",
            filename="paper.pdf",
            size_bytes=len(content),
        )
        document_id = document.id
        project_id = project.id
        await session.commit()

    async with session_factory() as session:
        await _build_pipeline(session, settings).run(
            document_id, project_id=project_id, content=content
        )
        await session.commit()

    async with session_factory() as session:
        refreshed = await session.get(Document, document_id)
        assert refreshed is not None
        assert refreshed.status == DocumentStatus.READY
        assert refreshed.metadata_["page_count"] == 3

        chunk_count = await session.scalar(
            select(func.count())
            .select_from(DocumentChunk)
            .where(DocumentChunk.document_id == document_id)
        )
        assert chunk_count is not None and chunk_count > 0

        # Every chunk must carry an embedding of the configured dimensionality
        # (design FR-C7) and the denormalised project_id (design §7.5).
        chunks = (
            (
                await session.execute(
                    select(DocumentChunk).where(DocumentChunk.document_id == document_id)
                )
            )
            .scalars()
            .all()
        )
        for chunk in chunks:
            assert chunk.embedding is not None
            assert len(chunk.embedding) == settings.embedding_dimensions
            assert chunk.project_id == project_id
            assert "page_start" in chunk.source_ref


async def test_docx_ingestion_records_section_provenance(
    db_engine: AsyncEngine, settings: Settings
) -> None:
    session_factory = async_sessionmaker(bind=db_engine, expire_on_commit=False)
    content = _docx_bytes()
    mime = "application/vnd.openxmlformats-officedocument.wordprocessingml.document"

    async with session_factory() as session:
        project = await _seed_project(session)
        document = await _create_pending_document(
            session, project, mime_type=mime, filename="paper.docx", size_bytes=len(content)
        )
        document_id = document.id
        project_id = project.id
        await session.commit()

    async with session_factory() as session:
        await _build_pipeline(session, settings).run(
            document_id, project_id=project_id, content=content
        )
        await session.commit()

    async with session_factory() as session:
        refreshed = await session.get(Document, document_id)
        assert refreshed is not None
        assert refreshed.status == DocumentStatus.READY

        chunks = (
            (
                await session.execute(
                    select(DocumentChunk).where(DocumentChunk.document_id == document_id)
                )
            )
            .scalars()
            .all()
        )
        sections = {chunk.source_ref.get("section") for chunk in chunks}
        assert "Introduction" in sections or "Conclusion" in sections


async def test_corrupt_pdf_marks_document_failed(
    db_engine: AsyncEngine, settings: Settings
) -> None:
    session_factory = async_sessionmaker(bind=db_engine, expire_on_commit=False)
    content = b"this is definitely not a pdf"

    async with session_factory() as session:
        project = await _seed_project(session)
        document = await _create_pending_document(
            session,
            project,
            mime_type="application/pdf",
            filename="broken.pdf",
            size_bytes=len(content),
        )
        document_id = document.id
        project_id = project.id
        await session.commit()

    async with session_factory() as session:
        await _build_pipeline(session, settings).run(
            document_id, project_id=project_id, content=content
        )
        await session.commit()

    async with session_factory() as session:
        refreshed = await session.get(Document, document_id)
        assert refreshed is not None
        assert refreshed.status == DocumentStatus.FAILED
        assert refreshed.metadata_["failure_reason"] is not None

        chunk_count = await session.scalar(
            select(func.count())
            .select_from(DocumentChunk)
            .where(DocumentChunk.document_id == document_id)
        )
        assert chunk_count == 0


async def test_reingestion_replaces_chunks_without_duplicating(
    db_engine: AsyncEngine, settings: Settings
) -> None:
    """Design §8.1: re-running "deletes and rebuilds its chunks rather than
    appending, so a retry after a crash cannot produce duplicates"."""
    session_factory = async_sessionmaker(bind=db_engine, expire_on_commit=False)
    content = _multi_page_pdf_bytes()

    async with session_factory() as session:
        project = await _seed_project(session)
        document = await _create_pending_document(
            session,
            project,
            mime_type="application/pdf",
            filename="paper.pdf",
            size_bytes=len(content),
        )
        document_id = document.id
        project_id = project.id
        await session.commit()

    async def run_once() -> None:
        async with session_factory() as session:
            await _build_pipeline(session, settings).run(
                document_id, project_id=project_id, content=content
            )
            await session.commit()

    await run_once()
    async with session_factory() as session:
        first_count = await session.scalar(
            select(func.count())
            .select_from(DocumentChunk)
            .where(DocumentChunk.document_id == document_id)
        )

    await run_once()
    async with session_factory() as session:
        second_count = await session.scalar(
            select(func.count())
            .select_from(DocumentChunk)
            .where(DocumentChunk.document_id == document_id)
        )

    assert first_count == second_count
