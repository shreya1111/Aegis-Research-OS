"""Unit tests for DOCX parsing (design §8.2, FR-C3)."""

from __future__ import annotations

import io

import pytest
from docx import Document as DocxDocument

from app.services.ingestion.errors import ExtractionError
from app.services.ingestion.parsers.docx import parse_docx


def _build_docx(*, title: str | None = None, author: str | None = None) -> bytes:
    document = DocxDocument()
    if title:
        document.core_properties.title = title
    if author:
        document.core_properties.author = author

    document.add_heading("Introduction", level=1)
    document.add_paragraph("This is the introduction paragraph.")
    document.add_heading("Methodology", level=1)
    document.add_paragraph("This is the methodology paragraph.")

    buffer = io.BytesIO()
    document.save(buffer)
    return buffer.getvalue()


@pytest.mark.unit
class TestParseDocx:
    def test_extracts_paragraph_text(self) -> None:
        extracted = parse_docx(_build_docx())
        full_text = extracted.full_text()
        assert "introduction paragraph" in full_text
        assert "methodology paragraph" in full_text

    def test_extracts_heading_structure_as_sections(self) -> None:
        extracted = parse_docx(_build_docx())
        sections = {page.section for page in extracted.pages if page.section}
        assert "Introduction" in sections
        assert "Methodology" in sections

    def test_extracts_title_metadata(self) -> None:
        extracted = parse_docx(_build_docx(title="My Thesis"))
        assert extracted.title == "My Thesis"

    def test_extracts_author_metadata(self) -> None:
        extracted = parse_docx(_build_docx(author="Ada Lovelace"))
        assert extracted.authors == ["Ada Lovelace"]

    def test_title_is_none_when_absent(self) -> None:
        extracted = parse_docx(_build_docx())
        assert extracted.title is None

    def test_pages_never_used_ocr(self) -> None:
        extracted = parse_docx(_build_docx())
        assert extracted.used_ocr is False

    def test_raises_extraction_error_for_a_corrupt_file(self) -> None:
        with pytest.raises(ExtractionError):
            parse_docx(b"this is not a docx file, just plain bytes")

    def test_raises_extraction_error_for_an_empty_document(self) -> None:
        document = DocxDocument()
        buffer = io.BytesIO()
        document.save(buffer)

        with pytest.raises(ExtractionError):
            parse_docx(buffer.getvalue())
